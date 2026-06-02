import openpyxl
from io import BytesIO

from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.views.generic import ListView, CreateView, UpdateView, DeleteView, DetailView, View
from django.urls import reverse_lazy
from django.contrib import messages
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.models import User
from django.template.loader import get_template
from django.utils.decorators import method_decorator
from django.contrib.auth.decorators import login_required

import xhtml2pdf.pisa as pisa

from datetime import date

from .models import Member, Staff, BaptismCertificate
from .forms import MemberForm, StaffForm
from accounts.decorators import admin_only, pastor_or_admin
from accounts.models import ChurchInfo


def render_to_pdf(template_src, context_dict={}):
    template = get_template(template_src)
    html = template.render(context_dict)
    result = BytesIO()
    pdf = pisa.pisaDocument(BytesIO(html.encode("UTF-8")), result)
    if not pdf.err:
        return HttpResponse(result.getvalue(), content_type='application/pdf')
    return None


class AdminRequiredMixin(UserPassesTestMixin):
    def test_func(self):
        return self.request.user.is_authenticated and self.request.user.profile.role == 'admin'

    def handle_no_permission(self):
        messages.error(self.request, 'Admin access required.')
        return redirect('dashboard')


class PastorOrAdminViewMixin(UserPassesTestMixin):
    def test_func(self):
        return self.request.user.is_authenticated and self.request.user.profile.role in ['admin', 'pastor']

    def handle_no_permission(self):
        messages.error(self.request, 'Pastor or Admin access required.')
        return redirect('dashboard')


class StaffOrAboveMixin(UserPassesTestMixin):
    def test_func(self):
        return self.request.user.is_authenticated and self.request.user.profile.role in ['admin', 'pastor', 'staff']

    def handle_no_permission(self):
        messages.error(self.request, 'You do not have permission to access this page.')
        return redirect('dashboard')


# ---- Member Views ----


@method_decorator(login_required, name='dispatch')
class MemberListView(LoginRequiredMixin, StaffOrAboveMixin, ListView):
    model = Member
    template_name = 'members/member_list.html'
    context_object_name = 'members'
    paginate_by = 25

    def get_queryset(self):
        qs = super().get_queryset()
        query = self.request.GET.get('q')
        voter = self.request.GET.get('voter')
        if query:
            qs = qs.filter(
                Q(first_name__icontains=query) |
                Q(last_name__icontains=query) |
                Q(membership_id__icontains=query) |
                Q(phone__icontains=query) |
                Q(email__icontains=query) |
                Q(voter_id__icontains=query)
            )
        if voter == 'yes':
            qs = qs.filter(is_voter=True)
        elif voter == 'no':
            qs = qs.filter(is_voter=False)
        return qs

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx['search_query'] = self.request.GET.get('q', '')
        ctx['voter_filter'] = self.request.GET.get('voter', '')
        return ctx


@method_decorator(login_required, name='dispatch')
class MemberDetailView(LoginRequiredMixin, StaffOrAboveMixin, DetailView):
    model = Member
    template_name = 'members/member_detail.html'
    context_object_name = 'member'


@method_decorator(login_required, name='dispatch')
class MemberCreateView(LoginRequiredMixin, AdminRequiredMixin, CreateView):
    model = Member
    form_class = MemberForm
    template_name = 'members/member_form.html'
    success_url = reverse_lazy('members:member_list')

    def form_valid(self, form):
        response = super().form_valid(form)
        messages.success(self.request, f'Member {self.object.full_name()} created successfully.')
        return response

    def form_invalid(self, form):
        messages.error(self.request, 'Please correct the errors below.')
        return super().form_invalid(form)


@method_decorator(login_required, name='dispatch')
class MemberUpdateView(LoginRequiredMixin, AdminRequiredMixin, UpdateView):
    model = Member
    form_class = MemberForm
    template_name = 'members/member_form.html'
    context_object_name = 'member'

    def get_success_url(self):
        return reverse_lazy('members:member_detail', kwargs={'pk': self.object.pk})

    def form_valid(self, form):
        response = super().form_valid(form)
        messages.success(self.request, f'Member {self.object.full_name()} updated successfully.')
        return response

    def form_invalid(self, form):
        messages.error(self.request, 'Please correct the errors below.')
        return super().form_invalid(form)


@method_decorator(login_required, name='dispatch')
class MemberDeleteView(LoginRequiredMixin, AdminRequiredMixin, DeleteView):
    model = Member
    template_name = 'members/member_confirm_delete.html'
    context_object_name = 'member'
    success_url = reverse_lazy('members:member_list')

    def form_valid(self, form):
        messages.success(self.request, f'Member {self.object.full_name()} deleted successfully.')
        return super().form_valid(form)


@method_decorator(login_required, name='dispatch')
class MemberPDFView(LoginRequiredMixin, StaffOrAboveMixin, DetailView):
    model = Member
    template_name = 'member_pdf.html'

    def render_to_response(self, context, **response_kwargs):
        member = self.object
        today = date.today()
        age = today.year - member.date_of_birth.year - ((today.month, today.day) < (member.date_of_birth.month, member.date_of_birth.day)) if member.date_of_birth else ''
        context = {
            'member': member,
            'age': age,
            'generated_date': today.strftime('%B %d, %Y'),
            'church_info': ChurchInfo.get_info(),
            'photo_path': member.photo.path if member.photo else None,
        }
        pdf = render_to_pdf('member_pdf.html', context)
        if pdf:
            return pdf
        return HttpResponse('PDF generation failed.', status=500)


@method_decorator(login_required, name='dispatch')
class MemberIDCardPDFView(LoginRequiredMixin, StaffOrAboveMixin, DetailView):
    model = Member

    def render_to_response(self, context, **response_kwargs):
        member = self.object
        today = date.today()
        context = {
            'member': member,
            'generated_date': today.strftime('%B %d, %Y'),
            'valid_until': date(today.year + 1, today.month, today.day).strftime('%B %d, %Y'),
            'church_info': ChurchInfo.get_info(),
            'photo_path': member.photo.path if member.photo else None,
        }
        pdf = render_to_pdf('id_card_pdf.html', context)
        if pdf:
            return pdf
        return HttpResponse('PDF generation failed.', status=500)


# ---- Staff Views ----


@method_decorator(login_required, name='dispatch')
class StaffListView(LoginRequiredMixin, PastorOrAdminViewMixin, ListView):
    model = Staff
    template_name = 'members/staff_list.html'
    context_object_name = 'staff_list'
    paginate_by = 25

    def get_queryset(self):
        qs = super().get_queryset()
        query = self.request.GET.get('q')
        if query:
            qs = qs.filter(
                Q(first_name__icontains=query) |
                Q(last_name__icontains=query) |
                Q(staff_id__icontains=query) |
                Q(role__icontains=query) |
                Q(email__icontains=query)
            )
        return qs

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx['search_query'] = self.request.GET.get('q', '')
        return ctx


@method_decorator(login_required, name='dispatch')
class StaffDetailView(LoginRequiredMixin, PastorOrAdminViewMixin, DetailView):
    model = Staff
    template_name = 'members/staff_detail.html'
    context_object_name = 'staff'


@method_decorator(login_required, name='dispatch')
class StaffCreateView(LoginRequiredMixin, AdminRequiredMixin, CreateView):
    model = Staff
    form_class = StaffForm
    template_name = 'members/staff_form.html'
    success_url = reverse_lazy('members:staff_list')

    def form_valid(self, form):
        response = super().form_valid(form)
        messages.success(self.request, f'Staff {self.object.full_name()} created successfully.')
        return response

    def form_invalid(self, form):
        messages.error(self.request, 'Please correct the errors below.')
        return super().form_invalid(form)


@method_decorator(login_required, name='dispatch')
class StaffUpdateView(LoginRequiredMixin, AdminRequiredMixin, UpdateView):
    model = Staff
    form_class = StaffForm
    template_name = 'members/staff_form.html'
    context_object_name = 'staff'

    def get_success_url(self):
        return reverse_lazy('members:staff_detail', kwargs={'pk': self.object.pk})

    def form_valid(self, form):
        response = super().form_valid(form)
        messages.success(self.request, f'Staff {self.object.full_name()} updated successfully.')
        return response

    def form_invalid(self, form):
        messages.error(self.request, 'Please correct the errors below.')
        return super().form_invalid(form)


@method_decorator(login_required, name='dispatch')
class StaffDeleteView(LoginRequiredMixin, AdminRequiredMixin, DeleteView):
    model = Staff
    template_name = 'members/staff_confirm_delete.html'
    context_object_name = 'staff'
    success_url = reverse_lazy('members:staff_list')

    def form_valid(self, form):
        messages.success(self.request, f'Staff {self.object.full_name()} deleted successfully.')
        return super().form_valid(form)


@method_decorator(login_required, name='dispatch')
class MemberImportView(LoginRequiredMixin, AdminRequiredMixin, View):
    template_name = 'members/member_import.html'

    def get(self, request):
        return render(request, self.template_name)

    def post(self, request):
        excel_file = request.FILES.get('excel_file')
        if not excel_file:
            messages.error(request, 'Please select an Excel file.')
            return render(request, self.template_name)

        if not excel_file.name.endswith(('.xlsx', '.xls')):
            messages.error(request, 'Please upload a .xlsx or .xls file.')
            return render(request, self.template_name)

        try:
            wb = openpyxl.load_workbook(excel_file, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
        except Exception as e:
            messages.error(request, f'Could not read Excel file: {e}')
            return render(request, self.template_name)

        if len(rows) < 2:
            messages.error(request, 'Excel file must have a header row and at least one data row.')
            return render(request, self.template_name)

        headers = [str(h).strip().lower().replace(' ', '_') if h else '' for h in rows[0]]

        FIELD_MAP = {
            'first_name': 'first_name', 'firstname': 'first_name', 'first name': 'first_name',
            'last_name': 'last_name', 'lastname': 'last_name', 'last name': 'last_name',
            'gender': 'gender',
            'date_of_birth': 'date_of_birth', 'dob': 'date_of_birth', 'birth_date': 'date_of_birth', 'birthdate': 'date_of_birth',
            'phone': 'phone', 'mobile': 'phone', 'contact': 'phone', 'phone_number': 'phone',
            'email': 'email', 'e_mail': 'email',
            'address': 'address',
            'occupation': 'occupation',
            'marital_status': 'marital_status', 'maritalstatus': 'marital_status',
            'baptism_status': 'baptism_status', 'baptismstatus': 'baptism_status',
            'baptism_date': 'baptism_date', 'baptismdate': 'baptism_date',
            'joining_date': 'joining_date', 'joined_date': 'joining_date', 'join_date': 'joining_date',
            'emergency_contact_name': 'emergency_contact_name', 'emergencyname': 'emergency_contact_name',
            'emergency_contact_phone': 'emergency_contact_phone', 'emergencyphone': 'emergency_contact_phone',
            'emergency_contact_relation': 'emergency_contact_relation', 'emergencyrelation': 'emergency_contact_relation',
            'is_active': 'is_active', 'active': 'is_active',
            'is_voter': 'is_voter', 'voter': 'is_voter',
            'voter_id': 'voter_id', 'voterid': 'voter_id', 'epic': 'voter_id',
            'notes': 'notes',
        }

        column_map = {}
        for i, h in enumerate(headers):
            clean = h.strip().lower().replace(' ', '_')
            if clean in FIELD_MAP:
                column_map[i] = FIELD_MAP[clean]

        required = ['first_name', 'last_name', 'gender', 'date_of_birth', 'phone', 'address']
        missing = [f for f in required if f not in column_map.values()]
        if missing:
            messages.error(request, f'Missing required columns: {", ".join(missing)}')
            return render(request, self.template_name)

        GENDER_MAP = {'male': 'male', 'm': 'male', 'male': 'male', 'female': 'female', 'f': 'female', 'female': 'female', 'other': 'other', 'o': 'other'}
        MARITAL_MAP = {'single': 'single', 's': 'single', 'married': 'married', 'm': 'married', 'divorced': 'divorced', 'd': 'divorced', 'widowed': 'widowed', 'w': 'widowed'}
        BAPTISM_MAP = {'water': 'water', 'w': 'water', 'holy_spirit': 'holy_spirit', 'holy spirit': 'holy_spirit', 'hs': 'holy_spirit', 'spirit': 'holy_spirit', 'both': 'both', 'b': 'both', 'none': 'none', 'n': 'none'}
        BOOL_MAP = {'yes': True, 'y': True, 'true': True, '1': True, 'no': False, 'n': False, 'false': False, '0': False}

        imported = 0
        errors = []

        for row_idx, row in enumerate(rows[1:], start=2):
            try:
                data = {}
                for col_idx, field_name in column_map.items():
                    val = row[col_idx] if col_idx < len(row) else None
                    data[field_name] = val

                first_name = str(data.get('first_name', '') or '').strip()
                last_name = str(data.get('last_name', '') or '').strip()
                gender_raw = str(data.get('gender', '') or '').strip().lower()
                phone = str(data.get('phone', '') or '').strip()
                address = str(data.get('address', '') or '').strip()

                row_errors = []
                if not first_name:
                    row_errors.append('first_name is empty')
                if not last_name:
                    row_errors.append('last_name is empty')
                if not gender_raw:
                    row_errors.append('gender is empty')
                gender = GENDER_MAP.get(gender_raw)
                if not gender:
                    row_errors.append(f'invalid gender "{gender_raw}"')
                if not phone:
                    row_errors.append('phone is empty')
                if not address:
                    row_errors.append('address is empty')

                dob_raw = data.get('date_of_birth')
                from datetime import datetime
                if isinstance(dob_raw, datetime):
                    date_of_birth = dob_raw.date()
                elif isinstance(dob_raw, date):
                    date_of_birth = dob_raw
                elif dob_raw:
                    try:
                        date_of_birth = datetime.strptime(str(dob_raw).strip(), '%Y-%m-%d').date()
                    except ValueError:
                        try:
                            date_of_birth = datetime.strptime(str(dob_raw).strip(), '%d/%m/%Y').date()
                        except ValueError:
                            try:
                                date_of_birth = datetime.strptime(str(dob_raw).strip(), '%m/%d/%Y').date()
                            except ValueError:
                                row_errors.append(f'invalid date_of_birth "{dob_raw}"')
                                date_of_birth = None
                else:
                    row_errors.append('date_of_birth is empty')
                    date_of_birth = None

                if row_errors:
                    errors.append(f'Row {row_idx}: {"; ".join(row_errors)}')
                    continue

                baptism_date = None
                bd_raw = data.get('baptism_date')
                if bd_raw:
                    if isinstance(bd_raw, datetime):
                        baptism_date = bd_raw.date()
                    elif isinstance(bd_raw, date):
                        baptism_date = bd_raw
                    else:
                        try:
                            baptism_date = datetime.strptime(str(bd_raw).strip(), '%Y-%m-%d').date()
                        except ValueError:
                            try:
                                baptism_date = datetime.strptime(str(bd_raw).strip(), '%d/%m/%Y').date()
                            except ValueError:
                                try:
                                    baptism_date = datetime.strptime(str(bd_raw).strip(), '%m/%d/%Y').date()
                                except ValueError:
                                    pass

                joining_date = date.today()
                jd_raw = data.get('joining_date')
                if jd_raw:
                    if isinstance(jd_raw, datetime):
                        joining_date = jd_raw.date()
                    elif isinstance(jd_raw, date):
                        joining_date = jd_raw
                    else:
                        try:
                            joining_date = datetime.strptime(str(jd_raw).strip(), '%Y-%m-%d').date()
                        except ValueError:
                            try:
                                joining_date = datetime.strptime(str(jd_raw).strip(), '%d/%m/%Y').date()
                            except ValueError:
                                try:
                                    joining_date = datetime.strptime(str(jd_raw).strip(), '%m/%d/%Y').date()
                                except ValueError:
                                    pass

                marital_raw = str(data.get('marital_status', '') or '').strip().lower()
                marital_status = MARITAL_MAP.get(marital_raw, 'single')

                baptism_raw = str(data.get('baptism_status', '') or '').strip().lower()
                baptism_status = BAPTISM_MAP.get(baptism_raw, 'none')

                active_raw = str(data.get('is_active', '') or '').strip().lower()
                is_active = BOOL_MAP.get(active_raw, True)

                voter_raw = str(data.get('is_voter', '') or '').strip().lower()
                is_voter = BOOL_MAP.get(voter_raw, False)

                voter_id = str(data.get('voter_id', '') or '').strip()

                Member.objects.create(
                    first_name=first_name,
                    last_name=last_name,
                    gender=gender,
                    date_of_birth=date_of_birth,
                    phone=phone,
                    email=str(data.get('email', '') or '').strip(),
                    address=address,
                    occupation=str(data.get('occupation', '') or '').strip(),
                    marital_status=marital_status,
                    baptism_status=baptism_status,
                    baptism_date=baptism_date,
                    joining_date=joining_date,
                    emergency_contact_name=str(data.get('emergency_contact_name', '') or '').strip(),
                    emergency_contact_phone=str(data.get('emergency_contact_phone', '') or '').strip(),
                    emergency_contact_relation=str(data.get('emergency_contact_relation', '') or '').strip(),
                    is_active=is_active,
                    is_voter=is_voter,
                    voter_id=voter_id,
                    notes=str(data.get('notes', '') or '').strip(),
                )
                imported += 1

            except Exception as e:
                errors.append(f'Row {row_idx}: {e}')

        summary = f'<strong>Import complete.</strong> {imported} member(s) imported successfully.'
        if errors:
            summary += f' <strong>{len(errors)} error(s)</strong>:<br>' + '<br>'.join(errors)
        messages.success(request, summary)
        return render(request, self.template_name, {'imported': imported, 'errors': errors})


@method_decorator(login_required, name='dispatch')
class MemberImportSampleView(LoginRequiredMixin, StaffOrAboveMixin, View):

    def get(self, request):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Members'

        headers = [
            'first_name', 'last_name', 'gender', 'date_of_birth', 'phone',
            'email', 'address', 'occupation', 'marital_status', 'baptism_status',
            'baptism_date', 'joining_date', 'emergency_contact_name',
            'emergency_contact_phone', 'emergency_contact_relation',
            'is_active', 'is_voter', 'voter_id', 'notes',
        ]
        ws.append(headers)

        ws.append(['John', 'Doe', 'male', '1990-05-15', '9876543210',
                    'john@email.com', '123 Main St, City', 'Engineer', 'married', 'water',
                    '2020-01-12', '2023-01-01', 'Jane Doe', '9876543211', 'Spouse',
                    'yes', 'yes', 'ABC1234567', 'Regular attendee'])
        ws.append(['Mary', 'Smith', 'female', '1985-08-22', '9876543212',
                    'mary@email.com', '456 Oak Ave, Town', 'Teacher', 'married', 'both',
                    '2019-06-08', '2022-06-01', 'Tom Smith', '9876543213', 'Spouse',
                    'yes', 'no', '', ''])
        ws.append(['Samuel', 'Johnson', 'male', '2000-11-03', '9876543214',
                    '', '789 Pine Rd, Village', 'Student', 'single', 'holy_spirit',
                    '2023-03-20', '2024-09-15', '', '', '',
                    'yes', 'no', '', 'Youth member'])

        for col in ws.columns:
            max_len = max(len(str(c.value or '')) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, 30)

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="member_import_template.xlsx"'
        wb.save(response)
        return response


@method_decorator(login_required, name='dispatch')
class StaffPDFView(LoginRequiredMixin, PastorOrAdminViewMixin, DetailView):
    model = Staff

    def render_to_response(self, context, **response_kwargs):
        staff = self.object
        context = {'staff': staff}
        pdf = render_to_pdf('staff_pdf.html', context)
        if pdf:
            return pdf
        return HttpResponse('PDF generation failed.', status=500)


# ---- Baptism Certificate Views ----


@method_decorator(login_required, name='dispatch')
class BaptismCertificateCreateView(LoginRequiredMixin, AdminRequiredMixin, CreateView):
    model = BaptismCertificate
    template_name = 'members/baptism_certificate_form.html'
    fields = ['baptism_date', 'baptism_type', 'officiated_by', 'notes']

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx['member'] = get_object_or_404(Member, pk=self.kwargs['member_pk'])
        return ctx

    def form_valid(self, form):
        member = get_object_or_404(Member, pk=self.kwargs['member_pk'])
        form.instance.member = member
        response = super().form_valid(form)
        messages.success(self.request, f'Baptism certificate for {member.full_name()} created successfully.')
        return response

    def get_success_url(self):
        return reverse_lazy('members:member_detail', kwargs={'pk': self.kwargs['member_pk']})


@method_decorator(login_required, name='dispatch')
class BaptismCertificatePDFView(LoginRequiredMixin, StaffOrAboveMixin, DetailView):
    model = BaptismCertificate

    def render_to_response(self, context, **response_kwargs):
        cert = self.object
        today = date.today()
        context = {
            'certificate': cert,
            'church_info': ChurchInfo.get_info(),
            'generated_date': today.strftime('%B %d, %Y'),
        }
        pdf = render_to_pdf('baptism_certificate_pdf.html', context)
        if pdf:
            return pdf
        return HttpResponse('PDF generation failed.', status=500)


@method_decorator(login_required, name='dispatch')
class BaptismCertificateDeleteView(LoginRequiredMixin, AdminRequiredMixin, DeleteView):
    model = BaptismCertificate

    def get_success_url(self):
        member_pk = self.object.member.pk
        messages.success(self.request, 'Baptism certificate deleted successfully.')
        return reverse_lazy('members:member_detail', kwargs={'pk': member_pk})
